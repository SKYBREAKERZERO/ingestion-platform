from __future__ import annotations

from datetime import date, datetime, time
from pathlib import Path
from typing import Any

import openpyxl

from app.loader.base_loader import BaseLoader
from app.model.block import (
    BlockType,
    DocumentBlock,
)
from app.model.document import Document
from app.model.page import Page


class ExcelLoaderError(RuntimeError):
    """
    XLSX 文件加载异常。
    """


class ExcelLoader(BaseLoader):
    """
    XLSX 原始数据加载器。

    负责：
        - 加载 .xlsx 文件
        - 按 Workbook 原始 Sheet 顺序读取
        - 每个 Sheet 生成一个逻辑 Page
        - 每个非空 Row 生成一个 TABLE Block
        - 保留 Cell 列位置
        - 保留 0 / False
        - 保留公式文本
        - 保存 Sheet / Row 元数据
        - 保存隐藏 Sheet 状态
        - 输出统一 Document

    不负责：
        - 删除隐藏 Sheet
        - 删除空 Sheet
        - 删除重复行
        - 删除合计行
        - 删除注释行
        - 表头识别
        - Chapter / Section 建模
        - Chunk
        - Token 统计
        - JSON / PostgreSQL 保存

    设计原则：
        Loader 尽量无损读取原始 XLSX 数据。

        Sheet 是否删除：
            -> SheetFilter

        Row 是否删除：
            -> RowFilter

        Loader 不提前执行具有业务语义的数据删除。
    """

    SUPPORTED_EXTENSIONS = {
        ".xlsx",
    }

    # ==================================================
    # Constructor
    # ==================================================

    def __init__(
        self,
        *,
        data_only: bool = False,
        read_only: bool = True,
        cell_separator: str = " | ",
        multiline_separator: str = " / ",
    ) -> None:

        if not cell_separator:
            raise ValueError(
                "cell_separator cannot be empty."
            )

        if not multiline_separator:
            raise ValueError(
                "multiline_separator cannot be empty."
            )

        self.data_only = data_only
        self.read_only = read_only

        self.cell_separator = (
            cell_separator
        )

        self.multiline_separator = (
            multiline_separator
        )

    # ==================================================
    # Public API
    # ==================================================

    def load(
        self,
        file_path: str,
    ) -> Document:

        path = self._validate_path(
            file_path
        )

        workbook = None

        try:

            workbook = (
                openpyxl.load_workbook(
                    filename=str(path),

                    # False:
                    #   保留 =SUM(A1:A3)
                    #
                    # True:
                    #   尝试读取 Excel 保存的缓存结果。
                    data_only=self.data_only,

                    # 大型 XLSX 建议使用 read_only，
                    # 避免一次性将全部 Cell 放进内存。
                    read_only=self.read_only,
                )
            )

            return self._build_document(
                path=path,
                workbook=workbook,
            )

        except ExcelLoaderError:
            raise

        except Exception as exc:

            raise ExcelLoaderError(
                "Failed to load XLSX file "
                f"'{path.name}': {exc}"
            ) from exc

        finally:

            if workbook is not None:

                try:
                    workbook.close()

                except Exception:
                    pass

    # ==================================================
    # Build Document
    # ==================================================

    def _build_document(
        self,
        *,
        path: Path,
        workbook,
    ) -> Document:

        blocks: list[
            DocumentBlock
        ] = []

        pages: list[
            Page
        ] = []

        sheet_records: list[
            dict[str, Any]
        ] = []

        global_order = 0

        total_row_count = 0
        non_empty_row_count = 0

        total_cell_count = 0
        non_empty_cell_count = 0

        empty_sheet_count = 0
        hidden_sheet_count = 0
        very_hidden_sheet_count = 0

        # ==============================================
        # Sheets
        # ==============================================

        for sheet_index, sheet in enumerate(
            workbook.worksheets
        ):

            sheet_number = (
                sheet_index + 1
            )

            sheet_name = str(
                sheet.title
            )

            sheet_state = str(
                getattr(
                    sheet,
                    "sheet_state",
                    "visible",
                )
                or "visible"
            )

            normalized_state = (
                sheet_state.casefold()
            )

            if normalized_state == "hidden":

                hidden_sheet_count += 1

            elif (
                normalized_state
                == "veryhidden"
            ):

                very_hidden_sheet_count += 1

            sheet_blocks: list[
                DocumentBlock
            ] = []

            sheet_text_lines: list[
                str
            ] = []

            sheet_total_row_count = 0
            sheet_non_empty_row_count = 0

            sheet_total_cell_count = 0
            sheet_non_empty_cell_count = 0

            # ==========================================
            # Rows
            # ==========================================

            for row_number, raw_row in enumerate(
                sheet.iter_rows(
                    values_only=True
                ),
                start=1,
            ):

                sheet_total_row_count += 1
                total_row_count += 1

                raw_values = list(
                    raw_row
                )

                # ======================================
                # Remove meaningless trailing cells
                # ======================================
                #
                # Example:
                #
                #     [None, "A", None, None]
                #
                # ->
                #
                #     [None, "A"]
                #
                # 注意：
                #   只删除最右侧空值。
                #
                #   开头和中间空 Cell 必须保留，
                #   否则列位置会发生变化。

                raw_values = (
                    self._trim_trailing_empty_cells(
                        raw_values
                    )
                )

                if not raw_values:
                    continue

                # ======================================
                # Normalize Cells
                # ======================================

                cells = [
                    self._normalize_cell_value(
                        value
                    )
                    for value
                    in raw_values
                ]

                current_non_empty_count = sum(
                    1
                    for cell
                    in cells
                    if cell
                )

                sheet_total_cell_count += len(
                    cells
                )

                total_cell_count += len(
                    cells
                )

                sheet_non_empty_cell_count += (
                    current_non_empty_count
                )

                non_empty_cell_count += (
                    current_non_empty_count
                )

                # ======================================
                # Empty Row
                # ======================================
                #
                # 完全空行不建立 Block。
                #
                # 注意：
                #   0
                #   False
                #
                # 已经会被正常转换成文本，
                # 不会在这里被误认为空。

                if not any(
                    cells
                ):
                    continue

                sheet_non_empty_row_count += 1
                non_empty_row_count += 1

                # ======================================
                # Row Text
                # ======================================

                row_text = (
                    self._build_row_text(
                        cells
                    )
                )

                # ======================================
                # TABLE Block
                # ======================================

                block = DocumentBlock(
                    block_type=BlockType.TABLE,

                    text=row_text,

                    order=global_order,

                    # 一个 Sheet 视为一个逻辑表格空间。
                    table_index=sheet_index,

                    # DocumentBlock 内使用 0-based。
                    row_index=(
                        row_number - 1
                    ),

                    cells=cells,

                    # 初始 Page 与 Sheet 一一对应。
                    page_number=sheet_number,

                    source="xlsx",

                    metadata={
                        "source": "xlsx",

                        "content_kind": (
                            "table"
                        ),

                        "sheet_index": (
                            sheet_index
                        ),

                        "sheet_number": (
                            sheet_number
                        ),

                        "sheet_name": (
                            sheet_name
                        ),

                        "sheet_state": (
                            sheet_state
                        ),

                        # Excel 行号使用 1-based。
                        "row_number": (
                            row_number
                        ),

                        "source_row_number": (
                            row_number
                        ),

                        "column_count": (
                            len(
                                cells
                            )
                        ),

                        "non_empty_cell_count": (
                            current_non_empty_count
                        ),

                        "logical_page_number": (
                            sheet_number
                        ),
                    },
                )

                blocks.append(
                    block
                )

                sheet_blocks.append(
                    block
                )

                sheet_text_lines.append(
                    row_text
                )

                global_order += 1

            # ==========================================
            # Sheet Page
            # ==========================================

            sheet_text = "\n".join(
                sheet_text_lines
            ).strip()

            if (
                sheet_non_empty_row_count
                == 0
            ):

                empty_sheet_count += 1

            # 即使 Sheet 为空，也建立 Page。
            #
            # 原因：
            #   SheetFilter 需要有机会决定
            #   是否删除这个 Sheet。

            pages.append(
                Page(
                    page_number=(
                        sheet_number
                    ),
                    text=sheet_text,
                )
            )

            # ==========================================
            # Sheet Metadata
            # ==========================================

            sheet_records.append(
                {
                    "sheet_index": (
                        sheet_index
                    ),

                    "sheet_number": (
                        sheet_number
                    ),

                    "sheet_name": (
                        sheet_name
                    ),

                    "sheet_state": (
                        sheet_state
                    ),

                    "logical_page_number": (
                        sheet_number
                    ),

                    "status": (
                        "EMPTY"
                        if (
                            sheet_non_empty_row_count
                            == 0
                        )
                        else "SUCCESS"
                    ),

                    "total_row_count": (
                        sheet_total_row_count
                    ),

                    "non_empty_row_count": (
                        sheet_non_empty_row_count
                    ),

                    "total_cell_count": (
                        sheet_total_cell_count
                    ),

                    "non_empty_cell_count": (
                        sheet_non_empty_cell_count
                    ),

                    "block_count": (
                        len(
                            sheet_blocks
                        )
                    ),

                    "character_count": (
                        len(
                            sheet_text
                        )
                    ),
                }
            )

        # ==============================================
        # Document Content Statistics
        # ==============================================

        character_count = sum(
            len(
                page.text
                or ""
            )
            for page
            in pages
        )

        # ==============================================
        # Metadata
        # ==============================================

        metadata: dict[
            str,
            Any,
        ] = {
            "source_format": (
                "xlsx"
            ),

            "loader": (
                "ExcelLoader"
            ),

            "loader_status": (
                "SUCCESS"
            ),

            "sheet_count": (
                len(
                    workbook.worksheets
                )
            ),

            "processed_sheet_count": (
                len(
                    workbook.worksheets
                )
            ),

            "empty_sheet_count": (
                empty_sheet_count
            ),

            "hidden_sheet_count": (
                hidden_sheet_count
            ),

            "very_hidden_sheet_count": (
                very_hidden_sheet_count
            ),

            "total_row_count": (
                total_row_count
            ),

            "non_empty_row_count": (
                non_empty_row_count
            ),

            "total_cell_count": (
                total_cell_count
            ),

            "non_empty_cell_count": (
                non_empty_cell_count
            ),

            "block_count": (
                len(
                    blocks
                )
            ),

            "page_count": (
                len(
                    pages
                )
            ),

            "character_count": (
                character_count
            ),

            "data_only": (
                self.data_only
            ),

            "read_only": (
                self.read_only
            ),

            "sheets": (
                sheet_records
            ),
        }

        # ==============================================
        # Document
        # ==============================================

        return Document(
            file_name=path.name,

            file_type="xlsx",

            pages=pages,

            blocks=blocks,

            chapters=[],

            sections=[],

            contents=[],

            metadata=metadata,
        )

    # ==================================================
    # Path Validation
    # ==================================================

    @classmethod
    def _validate_path(
        cls,
        file_path: str,
    ) -> Path:

        if not file_path:

            raise ValueError(
                "file_path cannot be empty."
            )

        path = Path(
            file_path
        ).expanduser()

        if not path.exists():

            raise FileNotFoundError(
                "XLSX file not found: "
                f"{path}"
            )

        if not path.is_file():

            raise IsADirectoryError(
                "Input path is not a file: "
                f"{path}"
            )

        # Excel 临时锁文件。
        if path.name.startswith(
            "~$"
        ):

            raise ValueError(
                "Temporary Excel file "
                "is not supported: "
                f"{path.name}"
            )

        if (
            path.suffix.lower()
            not in cls.SUPPORTED_EXTENSIONS
        ):

            raise ValueError(
                "ExcelLoader only accepts "
                ".xlsx files. "
                f"Received: "
                f"{path.suffix or '<no extension>'}"
            )

        return path

    # ==================================================
    # Cell Normalize
    # ==================================================

    def _normalize_cell_value(
        self,
        value: Any,
    ) -> str:
        """
        将 Excel Cell 转换成稳定文本。

        重点：

            None
                -> ""

            0
                -> "0"

            False
                -> "FALSE"

            True
                -> "TRUE"

            datetime/date/time
                -> ISO 格式

            Formula
                -> 原始公式字符串
                   （data_only=False 时）
        """

        if value is None:

            return ""

        # ==============================================
        # Boolean
        # ==============================================

        if isinstance(
            value,
            bool,
        ):

            return (
                "TRUE"
                if value
                else "FALSE"
            )

        # ==============================================
        # Date Time
        # ==============================================

        if isinstance(
            value,
            datetime,
        ):

            return value.isoformat(
                sep=" "
            )

        if isinstance(
            value,
            date,
        ):

            return value.isoformat()

        if isinstance(
            value,
            time,
        ):

            return value.isoformat()

        # ==============================================
        # String / Number / Formula / Error
        # ==============================================

        normalized = str(
            value
        )

        normalized = (
            normalized.replace(
                "\u3000",
                " ",
            )
        )

        normalized = (
            normalized.replace(
                "\xa0",
                " ",
            )
        )

        # ==============================================
        # Zero Width Characters
        # ==============================================

        for character in (
            "\u200b",
            "\u200c",
            "\u200d",
            "\u2060",
            "\ufeff",
        ):

            normalized = (
                normalized.replace(
                    character,
                    "",
                )
            )

        # ==============================================
        # Newline
        # ==============================================

        normalized = (
            normalized.replace(
                "\r\n",
                "\n",
            )
        )

        normalized = (
            normalized.replace(
                "\r",
                "\n",
            )
        )

        lines = [
            " ".join(
                line.split()
            )
            for line
            in normalized.splitlines()
            if line.strip()
        ]

        return self.multiline_separator.join(
            lines
        ).strip()

    # ==================================================
    # Trim Trailing Empty Cells
    # ==================================================

    @staticmethod
    def _trim_trailing_empty_cells(
        values: list[Any],
    ) -> list[Any]:
        """
        只删除行尾没有值的 Cell。

        Example:

            [
                None,
                "A",
                None,
                "B",
                None,
                None,
            ]

        ->

            [
                None,
                "A",
                None,
                "B",
            ]

        注意：

            开头的 None 保留。

            中间的 None 保留。

        因此不会导致实际数据列发生左移。
        """

        if not values:
            return []

        end = len(
            values
        )

        while (
            end > 0
            and values[
                end - 1
            ]
            is None
        ):

            end -= 1

        return values[
            :end
        ]

    # ==================================================
    # Row Text
    # ==================================================

    def _build_row_text(
        self,
        cells: list[str],
    ) -> str:
        """
        将 Row 转换成可检索文本。

        Example:

            [
                "",
                "Toyota",
                "",
                "100",
            ]

        ->

            | Toyota |  | 100

        cells 数组本身仍然保留完整列位置。
        """

        return self.cell_separator.join(
            cells
        ).strip()

    # ==================================================
    # Compatibility
    # ==================================================

    @staticmethod
    def _is_empty_value(
        value: Any,
    ) -> bool:
        """
        判断原始 Excel Cell 是否为空。

        这里只把：

            None

        视为空值。

        不把：

            0
            False

        视为空值。
        """

        return value is None


# ======================================================
# Compatibility Alias
# ======================================================
#
# 如果你的 Pipeline 当前使用：
#
#     from app.loader.excel_loader import XLSXLoader
#
# 或以后希望统一命名为 XLSXLoader，
# 不需要再维护两套 Loader。

XLSXLoader = ExcelLoader